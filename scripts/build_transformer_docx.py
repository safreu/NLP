"""Build the verified dual-decoder experiment report as a styled DOCX."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "transformer_experiments"
OUTPUT = ROOT / "docs" / "Dual_Decoder_Transformer_Experiments.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 39, 56)
MUTED = RGBColor(90, 100, 112)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
WHITE = RGBColor(255, 255, 255)


def set_font(run, size=11, color=INK, bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(indent_dxa))
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(widths_dxa[index]))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def add_header_footer(section) -> None:
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run("NLP Project  |  Transformer Experiments"), 9, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    set_font(footer.add_run("Dual-Decoder Transformer  •  "), 9, MUTED)
    add_page_field(footer)


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(3)
    set_font(kicker.add_run("CONTROLLED NLP EXPERIMENT REPORT"), 10, BLUE, True)
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("Dual-Decoder Transformer for Text Simplification"), 24, INK, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_font(
        subtitle.add_run(
            "Shared representation learning, reconstruction loss, and controlled internal ablations"
        ),
        13,
        MUTED,
    )
    metadata = [
        ("Dataset", "WikiLarge: 2,000 train / 200 validation / 191 test"),
        ("Controlled setup", "5 epochs • batch 8 • seed 42 • embedding 256"),
        ("Execution", "Completed CPU runs; server-ready Slurm and fallback scripts"),
        ("Final dual model", "ReLU • learned positions • 16 heads • reconstruction weight 1.00"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_font(paragraph.add_run(f"{label}: "), 10.5, INK, True)
        set_font(paragraph.add_run(value), 10.5, INK)
    doc.add_paragraph()


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    shade_cell(table.cell(0, 0), CALLOUT_FILL)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    set_font(paragraph.add_run(f"{label}: "), 10.5, DARK_BLUE, True)
    set_font(paragraph.add_run(text), 10.5, INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_configuration_table(doc: Document) -> None:
    rows = [
        ("E1", "Corrected", "0", "ReLU", "Learned", "8"),
        ("E2", "Dual", "0.25", "ReLU", "Learned", "8"),
        ("E3", "Dual", "0.50", "ReLU", "Learned", "8"),
        ("E4", "Dual", "1.00", "ReLU", "Learned", "8"),
        ("E5", "Dual", "1.00", "SwiGLU", "Learned", "8"),
        ("E6", "Dual", "1.00", "ReLU", "RoPE", "8"),
        ("E7", "Dual", "1.00", "ReLU", "Learned", "4"),
        ("E8", "Dual", "1.00", "ReLU", "Learned", "16"),
        ("E9", "Dual", "1.00", "ReLU", "Learned", "16"),
    ]
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ("Run", "Architecture", "λ", "Activation", "Position", "Heads")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        shade_cell(cell, LIGHT_FILL)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(value), 9, INK, True)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(cell.paragraphs[0].add_run(value), 9, INK)
    set_table_geometry(table, [760, 1880, 740, 1880, 1880, 980])


def load_results() -> list[dict[str, str]]:
    with (RESULTS / "comparison" / "all_results.csv").open(encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def value(row, key, digits=3):
    raw = row.get(key)
    if raw in (None, ""):
        return "NA"
    try:
        return f"{float(raw):.{digits}f}"
    except ValueError:
        return raw


def add_results_table(doc: Document, rows: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    headers = (
        "Run",
        "Params M",
        "Val loss",
        "Test loss",
        "SARI",
        "BERT F1",
        "ROUGE-L",
        "FK grade",
        "Entity",
    )
    for cell, label in zip(table.rows[0].cells, headers, strict=True):
        shade_cell(cell, LIGHT_FILL)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(label), 7.8, INK, True)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        data = (
            row["experiment"].split("_")[0],
            f"{float(row['parameters']) / 1_000_000:.2f}",
            value(row, "validation_loss"),
            value(row, "test_loss"),
            value(row, "sari"),
            value(row, "bert_f1"),
            value(row, "rouge_l"),
            value(row, "flesch_kincaid_grade"),
            value(row, "entity_preservation"),
        )
        cells = table.add_row().cells
        for cell, item in zip(cells, data, strict=True):
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(cell.paragraphs[0].add_run(item), 7.7, INK)
    set_table_geometry(table, [850, 930, 970, 970, 850, 980, 980, 970, 860])


def add_qualitative_examples(doc: Document) -> None:
    with (RESULTS / "comparison" / "qualitative_comparison.csv").open(encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    for position, index in enumerate((0, 2, 4, 5, 14), start=1):
        row = rows[index]
        heading = doc.add_paragraph(style="Heading 3")
        heading.add_run(
            f"Example {position} — {row['observed_error_categories'].replace(';', ', ')}"
        )
        fields = (
            ("Source", row["source"]),
            ("Reference", row["reference"]),
            ("Corrected baseline", row["corrected_baseline_output"]),
            ("Best dual / E9", row["final_combined_output"]),
            ("Reconstruction", row["best_dual_decoder_reconstruction"]),
        )
        for label, text in fields:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            set_font(paragraph.add_run(f"{label}: "), 9.5, DARK_BLUE, True)
            set_font(paragraph.add_run(text), 9.5, INK)
        observation = doc.add_paragraph()
        observation.paragraph_format.space_after = Pt(9)
        set_font(observation.add_run("Observation: "), 9.5, DARK_BLUE, True)
        set_font(
            observation.add_run(
                "The generated sentence does not reliably preserve the source meaning; the error "
                "labels come from the fixed deterministic comparison procedure."
            ),
            9.5,
            INK,
            italic=True,
        )


def build() -> Path:
    rows = load_results()
    selection = json.loads(
        (RESULTS / "comparison" / "final_configuration_selection.json").read_text()
    )
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    add_callout(
        doc,
        "Finding",
        "E9 is the best tested dual-decoder configuration under the balanced rule, but it does "
        "not improve consistently over the corrected single-decoder baseline.",
    )

    doc.add_heading("1. Objective", level=1)
    doc.add_paragraph(
        "The objective was to implement and evaluate a shared-encoder Transformer with a primary "
        "simplification decoder and an auxiliary source-reconstruction decoder, while controlling "
        "the reconstruction weight, feed-forward activation, positional encoding, and head count."
    )

    doc.add_heading("2. Baseline architecture and corrections", level=1)
    doc.add_paragraph(
        "The preserved historical model is a learned-position encoder-decoder Transformer. The "
        "corrected E1 baseline scales attention by the square root of each head dimension and "
        "combines causal and target-padding masks. Historical E0 results remain separately labeled."
    )
    add_bullets(
        doc,
        [
            "Source mask: (batch, 1, 1, source length), excluding source padding keys.",
            "Target mask: causal lower triangle AND target key-padding mask.",
            "Correct scaling: sqrt(head_dim), not sqrt(embed_size).",
        ],
    )

    doc.add_heading("3. Architectural modification", level=1)
    diagram = doc.add_paragraph()
    diagram.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        diagram.add_run(
            "Complex source  →  Shared encoder  →  Simplification decoder  →  Simple text\n"
        ),
        10.5,
        DARK_BLUE,
        True,
        name="Courier New",
    )
    set_font(
        diagram.add_run("                              ↘  Reconstruction decoder  →  Source text"),
        10.5,
        DARK_BLUE,
        True,
        name="Courier New",
    )
    doc.add_paragraph(
        "Both decoders attend to the same encoder states but have separate parameters. Normal "
        "inference executes only the simplification decoder; reconstruction is a training and "
        "diagnostic task."
    )

    doc.add_heading("4. Loss function and teacher forcing", level=1)
    equation = doc.add_paragraph()
    equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(equation.add_run("L = Lsimplification + λ · Lreconstruction"), 13, DARK_BLUE, True)
    doc.add_paragraph(
        "Simplification predicts target[1:] from target[:-1]. Reconstruction predicts source[1:] "
        "from source[:-1]. Both cross-entropy losses ignore padding tokens."
    )

    doc.add_heading("5. Internal modifications", level=1)
    add_bullets(
        doc,
        [
            "SwiGLU uses separate value and gate projections, SiLU on the gate, elementwise "
            "gating, and an output projection.",
            "RoPE rotates projected queries and keys before attention scores and removes learned "
            "absolute positions.",
            "At embedding 256, 4, 8, and 16 heads produce head dimensions 64, 32, and 16.",
            "All configurations validate divisibility; RoPE additionally requires an even head "
            "dimension.",
        ],
    )

    doc.add_heading("6. Controlled experiment plan", level=1)
    doc.add_paragraph(
        "All E1-E9 runs use the same split, vocabulary construction, tokenizer, seed, batch size, "
        "epoch count, optimizer, scheduler, checkpoint rule, and evaluation pipeline."
    )
    add_configuration_table(doc)

    doc.add_heading("7. Implementation and execution", level=1)
    doc.add_paragraph(
        "The CLI supports single runs, evaluation, generation, full dependency-ordered execution, "
        "resume, status, and comparison building. Every run stores configuration, environment, "
        "history, predictions, metrics, best/last checkpoints, logs, and status. Slurm is "
        "preferred; "
        "tmux, screen, and nohup are guarded fallbacks."
    )
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), CALLOUT_FILL)
    code._p.get_or_add_pPr().append(shade)
    set_font(
        code.add_run("bash scripts/slurm/submit_transformer_pipeline.sh"),
        9.5,
        INK,
        name="Courier New",
    )

    doc.add_heading("8. Results", level=1)
    doc.add_paragraph(
        "Lower loss and FK grade are better; higher SARI, BERTScore, ROUGE-L, and preservation are "
        "better. Weighted total losses are not directly comparable across reconstruction weights."
    )
    add_results_table(doc, rows)
    source_note = doc.add_paragraph()
    source_note.paragraph_format.space_before = Pt(4)
    source_note.paragraph_format.space_after = Pt(4)
    set_font(
        source_note.add_run(
            "Note. Number preservation is NA because the normalized WikiLarge test representation "
            "contains no numeric tokens. E0 is historical and unretrained."
        ),
        8.5,
        MUTED,
        italic=True,
    )

    doc.add_heading("9. Analysis", level=1)
    doc.add_heading("Reconstruction weight", level=2)
    doc.add_paragraph(
        "Lambda 1.00 won the weight sweep. It had the best SARI, BERTScore F1, readability, and "
        "validation simplification loss among E2-E4. However, all dual weights remained worse than "
        "E1 on BERTScore and readability."
    )
    doc.add_heading("SwiGLU and RoPE", level=2)
    doc.add_paragraph(
        "SwiGLU improved semantic overlap and readability over matched E4 but reduced SARI and "
        "added "
        "parameters. RoPE achieved the highest SARI yet sharply worsened BERTScore, ROUGE-L, token "
        "F1, and readability; it was not selected."
    )
    doc.add_heading("Head count and final selection", level=2)
    doc.add_paragraph(
        f"The balanced rule selected {selection['source_experiment']}. Sixteen heads produced the "
        "best dual BERTScore, ROUGE-L, BLEU, token F1, entity estimate, and readability, although "
        "four and eight heads produced better SARI. E9 retrained the selected configuration."
    )
    add_callout(
        doc,
        "Interpretation",
        "The dual architecture is not a consistent improvement over E1. E9 is the dual-ablation "
        "winner; E1 remains the stronger efficiency, BERTScore, and readability comparator.",
    )

    doc.add_heading("10. Qualitative examples", level=1)
    doc.add_paragraph(
        "The consolidated file uses the first 20 test examples in fixed order. The five cases "
        "below "
        "are displayed from that non-cherry-picked set."
    )
    add_qualitative_examples(doc)

    doc.add_heading("11. Limitations", level=1)
    add_bullets(
        doc,
        [
            "The controlled 2,000-example training set is too small for reliable word-level "
            "generation.",
            "Only one seed was run, so uncertainty is unknown.",
            "Whitespace tokenization and greedy decoding encourage generic-output collapse.",
            "Lowercasing weakens NER; numeric deletion prevents number-preservation measurement.",
            "The auxiliary decoder adds substantial parameters and training time.",
        ],
    )

    doc.add_heading("12. Conclusion", level=1)
    doc.add_paragraph(
        "The implementation demonstrates the intended Transformer modifications and a reproducible "
        "ablation pipeline. Lambda 1.00 was selected, SwiGLU and RoPE produced conflicting "
        "trade-offs, "
        "and 16 heads won the balanced dual-model rule. The experiment does not support an "
        "unconditional performance claim: E9 slightly improves several overlap and preservation "
        "metrics but remains behind E1 on BERTScore, readability, efficiency, and qualitative "
        "reliability. A larger multi-seed subword-tokenized server run is the appropriate next "
        "step."
    )

    for section in doc.sections:
        add_header_footer(section)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
